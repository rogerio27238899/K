"""Gera a apresentação comercial Nunes & Lucato -> mineradora (resíduo têxtil).
Uso: python3 gerar_apresentacao.py  (rodar a partir da raiz do repositório)
"""
import os
from PIL import Image, ImageOps, ImageEnhance
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARQ = os.path.join(ROOT, "arquivos")
OUT_DIR = os.path.join(ROOT, "apresentacao-mineradora")
ASSETS = os.path.join(OUT_DIR, "assets")
os.makedirs(ASSETS, exist_ok=True)

FOTOS = [os.path.join(ARQ, f"WhatsApp Image 2026-09-28 at 12.16.{s}.jpeg") for s in ("25", "26", "27")]

# ---------- paleta (extraída da apresentação Nunes & Lucato) ----------
DARK = RGBColor(0x0B, 0x0A, 0x1F)
NAVY = RGBColor(0x1E, 0x1D, 0x72)
ROYAL = RGBColor(0x34, 0x32, 0xC9)
BLUE = RGBColor(0x1E, 0x86, 0xC9)
GREEN = RGBColor(0x12, 0xA8, 0x7B)
CARD = RGBColor(0xF4, 0xF5, 0xFC)
BORDER = RGBColor(0xDD, 0xDE, 0xEE)
TEXT = RGBColor(0x1B, 0x1B, 0x2F)
GRAY = RGBColor(0x6B, 0x6B, 0x80)
LIGHT = RGBColor(0xC9, 0xCA, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Montserrat"
FOOTER_LABEL = "Resíduo Têxtil · Mineração"


# ---------- assets ----------
def make_logos():
    """Gera logo escura e branca com fundo transparente a partir do recorte do PDF da Nunes."""
    src = os.path.join(ASSETS, "logo_src.png")
    im = Image.open(src).convert("L")
    alpha = ImageOps.invert(im).point(lambda v: 0 if v < 25 else min(255, int(v * 1.3)))
    for name, color in (("logo_dark.png", (0x2A, 0x2A, 0x33)), ("logo_white.png", (255, 255, 255))):
        out = Image.new("RGBA", im.size, color + (0,))
        out.putalpha(alpha)
        out.save(os.path.join(ASSETS, name))


def duotone(path, out, size=(1400, 1620)):
    """Foto em duotone azul, como nas capas da Nunes."""
    im = Image.open(path).convert("L")
    im = ImageEnhance.Contrast(im).enhance(1.15)
    im = ImageOps.colorize(im, black=(0x0B, 0x0A, 0x3A), white=(0x9A, 0x9E, 0xF0), mid=(0x2E, 0x2C, 0xA8))
    im = ImageOps.fit(im, size, Image.LANCZOS)
    im.save(out, quality=90)
    return out


def fit(path, out, size):
    im = Image.open(path).convert("RGB")
    ImageOps.fit(im, size, Image.LANCZOS).save(out, quality=90)
    return out


# ---------- helpers ----------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SLIDE_NO = [1]  # capa conta como 01, igual à apresentação da Nunes


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    return s


def hline(slide, x, y, w, color=BORDER, weight=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def text(slide, x, y, w, h, runs, size=12, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=None, line_spacing=1.15):
    """runs: str, ou lista de parágrafos; cada parágrafo é str ou lista de (texto, bold[, color, size])."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line_spacing
        parts = p if isinstance(p, list) else [(p, False)]
        for part in parts:
            t, bold = part[0], part[1]
            r = para.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(part[3] if len(part) > 3 else size)
            f.bold = bold
            f.color.rgb = part[2] if len(part) > 2 else color
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def eyebrow(slide, x, y, label, color=ROYAL):
    hline(slide, x, y + 0.1, 0.32, color, 1)
    text(slide, x + 0.45, y, 6, 0.25, label, size=9, color=color, spacing=200)


def header(slide, label, title_runs, sub=None):
    eyebrow(slide, 0.83, 0.55, label)
    text(slide, 0.83, 0.85, 11.7, 0.7, [title_runs], size=28, color=TEXT)
    if sub:
        text(slide, 0.83, 1.6, 11.5, 0.5, sub, size=12, color=GRAY)


def footer(slide):
    SLIDE_NO[0] += 1
    hline(slide, 0.83, 6.78, 11.67)
    slide.shapes.add_picture(os.path.join(ASSETS, "logo_dark.png"), Inches(0.83), Inches(6.9), height=Inches(0.3))
    text(slide, 8.5, 6.95, 4.0, 0.25, f"{FOOTER_LABEL}   |   {SLIDE_NO[0]:02d}", size=8, color=GRAY,
         align=PP_ALIGN.RIGHT)


def white_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    return s


def dark_bg(slide, img):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK
    left = rect(slide, 0, 0, 8.3, 7.5, DARK)
    left.fill.gradient()
    left.fill.gradient_angle = 0
    stops = left.fill.gradient_stops
    stops[0].color.rgb = DARK
    stops[0].position = 0
    stops[1].color.rgb = NAVY
    stops[1].position = 1.0
    slide.shapes.add_picture(img, Inches(8.3), 0, Inches(5.033), Inches(7.5))
    # linhas diagonais sutis
    for i in range(14):
        ln = slide.shapes.add_connector(1, Inches(-2 + i * 0.75), Inches(7.5), Inches(1.5 + i * 0.75), Inches(0))
        ln.line.color.rgb = RGBColor(0x2A, 0x29, 0x70)
        ln.line.width = Pt(0.5)


ICONS = os.path.join(ASSETS, "icons")  # ícones de linha Lucide (licença ISC), renderizados nas cores da marca


def icon(slide, name, x, y, size=0.42, color=None):
    """Ícone de linha no estilo da apresentação original. color: ROYAL, BLUE, GREEN ou WHITE."""
    tone = {GREEN: "green", WHITE: "white", BLUE: "blue"}.get(color, "royal")
    slide.shapes.add_picture(os.path.join(ICONS, f"{name}_{tone}.png"), Inches(x), Inches(y), Inches(size), Inches(size))


def card(slide, x, y, w, h, bar, title, body, num=None, title_size=13, title_lines=1, ico=None):
    """title_lines: linhas reservadas ao título. Use 2 quando algum título da fileira quebra,
    para o texto não encostar no título e ficar alinhado em todos os cards.
    ico: com ícone, o card segue o estilo da original (fundo lilás, sem borda, ícone no topo)."""
    if ico:
        rect(slide, x, y, w, h, CARD)
        icon(slide, ico, x + 0.3, y + 0.32, 0.42, bar)
        ty = y + 0.32 + 0.42 + 0.25
    else:
        rect(slide, x, y, w, h, WHITE, BORDER)
        rect(slide, x, y, w, 0.06, bar)
        ty = y + 0.3
    if num:
        text(slide, x + 0.3, ty, 1, 0.4, num, size=18, color=bar)
        ty += 0.5
    title_h = 0.3 + 0.25 * (title_lines - 1)
    text(slide, x + 0.3, ty, w - 0.6, title_h, [[(title, True)]], size=title_size, color=TEXT, line_spacing=1.05)
    by = ty + title_h + 0.15
    text(slide, x + 0.3, by, w - 0.6, h - (by - y) - 0.2, body, size=10, color=GRAY)


# =====================================================================
make_logos()
cover_img = duotone(FOTOS[2], os.path.join(ASSETS, "capa_duotone.jpg"))
end_img = duotone(FOTOS[0], os.path.join(ASSETS, "fim_duotone.jpg"))
f1 = fit(FOTOS[0], os.path.join(ASSETS, "foto1.jpg"), (800, 1000))
f2 = fit(FOTOS[1], os.path.join(ASSETS, "foto2.jpg"), (800, 1000))
f3 = fit(FOTOS[2], os.path.join(ASSETS, "foto3.jpg"), (800, 1000))
f_wide = fit(FOTOS[1], os.path.join(ASSETS, "foto_wide.jpg"), (1400, 900))

# ---------- 1. CAPA ----------
s = prs.slides.add_slide(BLANK)
dark_bg(s, cover_img)
eyebrow(s, 0.83, 1.35, "Proposta comercial  ·  Setor de Mineração", color=LIGHT)
text(s, 0.83, 1.8, 7.2, 2.2, [[("Resíduo têxtil como ", False)], [("combustível alternativo", True)],
                             [("para a mineração", False)]], size=38, color=WHITE, line_spacing=1.1)
hline(s, 0.83, 4.3, 1.1, LIGHT, 1)
text(s, 0.83, 4.5, 6.8, 0.9, "Energia de alto poder calorífico, baixo cloro e destinação ambientalmente adequada, "
     "com rastreabilidade da coleta à entrega.", size=14, color=LIGHT)
text(s, 0.83, 5.95, 4, 0.25, "Proposta apresentada pela", size=9, color=LIGHT, spacing=200)
s.shapes.add_picture(os.path.join(ASSETS, "logo_white.png"), Inches(0.83), Inches(6.25), height=Inches(0.5))

# ---------- 2. QUEM SOMOS ----------
s = white_slide()
header(s, "Quem somos", [("Gestão ambiental especializada em ", False), ("resíduos têxteis", True)])
text(s, 0.83, 1.75, 6.3, 1.4, "A Nunes & Lucato atua na cadeia de resíduos têxteis com estrutura própria de "
     "transformação, equipe treinada e rastreabilidade documental completa, da coleta à destinação final "
     "ambientalmente adequada.", size=12, color=GRAY)
hline(s, 0.83, 3.25, 6.3)
text(s, 0.83, 3.4, 6.5, 0.4, [[("Do resíduo ao combustível", False, ROYAL), ("  ·  com respaldo documental", False, TEXT)]],
     size=14)
s.shapes.add_picture(f_wide, Inches(7.55), Inches(1.75), Inches(4.95), Inches(2.3))
cards = [(ROYAL, "truck", "Coleta e transporte", "Com emissão de MTR, dando respaldo legal à logística."),
         (ROYAL, "filter", "Triagem técnica", "Separação por composição e retirada de contaminantes."),
         (ROYAL, "scissors", "Picotagem", "Adequação da granulometria à especificação do forno do cliente."),
         (GREEN, "package-check", "Fornecimento", "Lotes padronizados, ensacados e rastreados até a entrega.")]
for i, (c, ic, t, b) in enumerate(cards):
    card(s, 0.83 + i * 2.97, 4.25, 2.75, 2.35, c, t, b, ico=ic)
footer(s)

# ---------- 3. O RESÍDUO ----------
s = white_slide()
header(s, "O material", [("Aparas têxteis pós-industriais, ", False), ("prontas para energia", True)])
for i, f in enumerate((f1, f2, f3)):
    s.shapes.add_picture(f, Inches(0.83 + i * 2.35), Inches(1.75), Inches(2.2), Inches(2.75))
text(s, 0.83, 4.65, 7, 0.3, "Registro fotográfico do material ensacado e estocado em lotes.", size=9, color=GRAY)
items = [("factory", "Origem", "Aparas e sobras de corte da indústria de confecção: material limpo, seco e sem uso prévio."),
         ("layers", "Composição", "Mistura de aparas de fibras naturais e sintéticas, padronizada por lote para manter o alto PCI."),
         ("scissors", "Forma", "Picotado em fragmentos, o que facilita dosagem e alimentação contínua."),
         ("package", "Embalagem", "Sacos de até 30 kg, com manuseio simples e estocagem organizada por lote.")]
y = 1.75
for ic, t, b in items:
    icon(s, ic, 8.1, y + 0.02, 0.4)
    text(s, 8.75, y, 3.75, 0.3, [[(t, True)]], size=12, color=TEXT)
    text(s, 8.75, y + 0.3, 3.75, 0.6, b, size=10, color=GRAY)
    y += 1.12
footer(s)

# ---------- 4. LAUDO ----------
s = white_slide()
header(s, "Resultados analíticos", [("Laudo laboratorial: ", False), ("alto poder calorífico e baixo cloro", True)],
       sub="Amostra de tecido triturado (matriz resíduo sólido), analisada por laboratório acreditado em 2026.")
big = [(ROYAL, "flame", "9.953,8", "kcal/kg", "Poder Calorífico Inferior (PCI, base seca)", "Referência mínima: ≥ 1.800 kcal/kg"),
       (BLUE, "flask-conical", "< 0,05", "%", "Cloro (base seca)", "Referência máxima: ≤ 1,0 %"),
       (GREEN, "droplets", "3,77", "%", "Umidade", "96,23 % de sólidos")]
for i, (c, ic, v, u, t, ref) in enumerate(big):
    x = 0.83 + i * 3.95
    rect(s, x, 2.3, 3.75, 2.75, CARD)
    rect(s, x, 2.3, 3.75, 0.06, c)
    icon(s, ic, x + 3.75 - 0.8, 2.62, 0.45, c)
    text(s, x + 0.35, 2.65, 3.2, 0.9, [[(v, True, c, 40), ("  " + u, False, c, 14)]])
    text(s, x + 0.35, 3.65, 3.1, 0.6, [[(t, True)]], size=12)
    text(s, x + 0.35, 4.3, 3.1, 0.4, ref, size=10, color=GRAY)
hline(s, 0.83, 5.35, 11.67)
text(s, 0.83, 5.5, 11.6, 0.4, [[("PCS: 10.040,0 kcal/kg  ·  ", False, TEXT), ("Conclusão do laudo: ", True, ROYAL),
                               ("os parâmetros de PCI e cloro atendem aos limites da Resolução SIMA nº 145/2021 (art. 5º).", False, TEXT)]],
     size=12)
text(s, 0.83, 6.0, 11.6, 0.4, "Resultado da amostra analisada. O alto PCI vem da mistura de fibras, e a composição é mantida em todos os lotes fornecidos.",
     size=9, color=GRAY)
footer(s)

# ---------- 5. COMPARATIVO PCI ----------
s = white_slide()
header(s, "Comparativo energético", [("Mais energia por quilo que os ", False), ("combustíveis tradicionais", True)],
       sub="Poder calorífico inferior (kcal/kg)")
bars = [("Resíduo têxtil Nunes & Lucato (laudo)", 9954, ROYAL, True),
        ("Óleo combustível", 9600, LIGHT, False),
        ("Coque de petróleo", 8000, LIGHT, False),
        ("Carvão mineral", 6000, LIGHT, False),
        ("Lenha / cavaco de madeira", 3000, LIGHT, False),
        ("Mínimo exigido para CDR (SIMA)", 1800, BORDER, False)]
maxw, x0, y = 7.3, 4.4, 2.3
for label, v, c, hl in bars:
    text(s, 0.83, y + 0.05, 3.45, 0.4, [[(label, hl)]], size=11, color=TEXT if hl else GRAY, align=PP_ALIGN.RIGHT)
    rect(s, x0, y, maxw * v / 10000, 0.45, c)
    text(s, x0 + maxw * v / 10000 + 0.15, y + 0.07, 1.2, 0.35, [[(f"{v:,}".replace(",", "."), hl)]], size=12,
         color=ROYAL if hl else GRAY)
    y += 0.68
text(s, 0.83, 6.3, 11.6, 0.35, "Combustíveis de comparação: valores típicos de referência, que variam conforme origem e "
     "qualidade. Resíduo têxtil: PCI em base seca conforme laudo.", size=9, color=GRAY)
footer(s)

# ---------- 6. EQUIVALÊNCIA ----------
s = white_slide()
header(s, "Equivalência energética", [("1 tonelada de resíduo têxtil ", True), ("substitui aproximadamente…", False)],
       sub="≈ 41,7 GJ de energia por tonelada (PCI 9.953,8 kcal/kg)")
eq = [(ROYAL, "fuel", "1,0 t", "de óleo combustível"), (BLUE, "flame", "1,2 t", "de coque de petróleo"),
      (GREEN, "pickaxe", "1,7 t", "de carvão mineral"), (ROYAL, "wind", "~1.150 m³", "de gás natural")]
for i, (c, ic, v, t) in enumerate(eq):
    x = 0.83 + i * 2.97
    rect(s, x, 2.3, 2.75, 2.6, CARD)
    rect(s, x, 2.3, 2.75, 0.06, c)
    icon(s, ic, x + 2.75 / 2 - 0.22, 2.6, 0.44, c)
    text(s, x, 3.2, 2.75, 0.8, [[(v, True, c, 32)]], align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 4.1, 2.35, 0.6, t, size=12, color=TEXT, align=PP_ALIGN.CENTER)
hline(s, 0.83, 5.15, 11.67)
text(s, 0.83, 5.35, 11.6, 0.5, [[("Menos combustível fóssil comprado, ", True, ROYAL),
                                ("com o mesmo aporte térmico no processo.", False, TEXT)]], size=16)
text(s, 0.83, 6.0, 11.6, 0.35, "Estimativa teórica por equivalência de PCI, com base nos valores típicos do slide anterior. "
     "Taxa real de substituição a validar em teste de queima.", size=9, color=GRAY)
footer(s)

# ---------- 7. APLICAÇÕES ----------
s = white_slide()
header(s, "Aplicações na mineração", [("Onde o resíduo têxtil ", False), ("gera valor na sua operação", True)])
apps = [(ROYAL, "circle-dot", "01", "Fornos de pelotização", "Substituição parcial do combustível sólido e do gás natural no endurecimento de pelotas."),
        (BLUE, "flame", "02", "Fornos de calcinação", "Aporte térmico em fornos de cal e calcinação de calcário e dolomita."),
        (GREEN, "wind", "03", "Secadores de minério", "Geração de calor para secagem de concentrados e minério úmido."),
        (ROYAL, "factory", "04", "Caldeiras e geradores", "Vapor e energia térmica em caldeiras adaptadas a combustível sólido.")]
for i, (c, ic, n, t, b) in enumerate(apps):
    x = 0.83 + i * 2.97
    card(s, x, 1.8, 2.75, 3.3, c, t, b, ico=ic)
    text(s, x + 2.75 - 0.8, 2.0, 0.5, 0.3, n, size=11, color=GRAY, align=PP_ALIGN.RIGHT)
hline(s, 0.83, 5.45, 11.67)
text(s, 0.83, 5.6, 11.6, 0.8, [[("Uso como substituição parcial, ", True, ROYAL),
                                ("iniciando com teste de queima e dosagem controlada, sujeito ao licenciamento ambiental "
                                 "da unidade consumidora.", False, GRAY)]], size=12)
footer(s)

# ---------- 8. BENEFÍCIOS ----------
s = white_slide()
header(s, "Benefícios", [("Vantagens para a ", False), ("mineradora", True)])
ben = [(ROYAL, "trending-down", "Redução de custo energético", "Energia de alto PCI a um custo competitivo frente aos combustíveis convencionais."),
       (ROYAL, "fuel", "Menos combustível fóssil", "Substituição parcial de coque, carvão, óleo ou gás no processo térmico."),
       (GREEN, "leaf", "Agenda ESG", "Economia circular e aterro zero, com indicadores para relatórios de sustentabilidade."),
       (ROYAL, "shield-check", "Baixo teor de cloro", "Menor risco de corrosão e incrustação em fornos e dutos."),
       (ROYAL, "droplets", "Baixa umidade", "Material seco (3,77 %), com energia útil elevada na queima."),
       (GREEN, "qr-code", "Rastreabilidade total", "MTR, certificado de destinação e controle por lote, com segurança jurídica.")]
for i, (c, ic, t, b) in enumerate(ben):
    col, row = i % 3, i // 3
    card(s, 0.83 + col * 3.95, 1.7 + row * 2.5, 3.75, 2.3, c, t, b, ico=ic)
footer(s)

# ---------- 9. FLUXO ----------
s = white_slide()
header(s, "Fluxo operacional", [("Como preparamos o ", False), ("combustível", True)])
steps = [("truck", "Coleta", "Recolhimento na indústria geradora com MTR"),
         ("inbox", "Recebimento", "Conferência, pesagem e registro"),
         ("filter", "Triagem", "Retirada de contaminantes e metais"),
         ("scissors", "Picotagem", "Granulometria sob especificação"),
         ("package", "Ensacamento", "Sacos de até 30 kg, por lote"),
         ("send", "Expedição", "Entrega programada com MTR")]
w = 1.8
for i, (ic, t, b) in enumerate(steps):
    x = 0.83 + i * 1.975
    c = GREEN if i == len(steps) - 1 else ROYAL
    circ = rect(s, x + w / 2 - 0.45, 2.1, 0.9, 0.9, WHITE, c, shape=MSO_SHAPE.OVAL)
    circ.line.width = Pt(1.5)
    icon(s, ic, x + w / 2 - 0.22, 2.33, 0.44, c)
    if i < len(steps) - 1:
        hline(s, x + w / 2 + 0.55, 2.55, 0.95, BORDER, 1.5)
    text(s, x, 3.15, w, 0.3, f"{i + 1:02d}", size=10, color=c, align=PP_ALIGN.CENTER)
    text(s, x, 3.42, w, 0.35, [[(t, True)]], size=12, align=PP_ALIGN.CENTER)
    text(s, x, 3.77, w, 0.8, b, size=10, color=GRAY, align=PP_ALIGN.CENTER)
rect(s, 0.83, 4.7, 11.67, 1.55, CARD)
rect(s, 0.83, 4.7, 0.06, 1.55, ROYAL)
text(s, 1.15, 4.9, 11.1, 1.2, [[("Estrutura própria de preparação", True, ROYAL)],
                               [("Área de armazenagem coberta, organizada por lotes, e picotagem própria para adequar o material "
                                 "à alimentação do forno do cliente, com fornecimento regular e programado.", False, GRAY)]],
     size=12, line_spacing=1.3)
footer(s)

# ---------- 9b. CAPACIDADE E LOGÍSTICA ----------
s = white_slide()
header(s, "Capacidade e logística", [("Fornecimento recorrente de ", False), ("1.000 toneladas por mês", True)],
       sub="Base operacional em São Paulo/SP (Belenzinho), com entrega programada na unidade do cliente.")
cap = [(ROYAL, "boxes", "1.000", "t/mês", "Volume disponível", "Fornecimento contínuo e programado"),
       (BLUE, "truck", "14", "t", "Por carga (truck)", "Sacos de até 30 kg, carga padronizada"),
       (GREEN, "calendar-check", "~72", "cargas/mês", "Entregas mensais", "Cerca de 3 cargas por dia útil"),
       (ROYAL, "zap", "~41.700", "GJ/mês", "Energia disponível", "≈ 9,95 bilhões de kcal/mês")]
for i, (c, ic, v, u, t, b) in enumerate(cap):
    x = 0.83 + i * 2.97
    rect(s, x, 2.2, 2.75, 2.8, CARD)
    rect(s, x, 2.2, 2.75, 0.06, c)
    icon(s, ic, x + 0.3, 2.48, 0.42, c)
    text(s, x + 0.3, 3.05, 2.4, 0.7, [[(v, True, c, 28), (" " + u, False, c, 11)]])
    text(s, x + 0.3, 3.85, 2.3, 0.4, [[(t, True)]], size=12)
    text(s, x + 0.3, 4.2, 2.3, 0.7, b, size=10, color=GRAY)
hline(s, 0.83, 5.25, 11.67)
text(s, 0.83, 5.45, 11.6, 0.5, [[("Volume recorrente e escala ", True, ROYAL),
                                ("para substituir parte relevante do combustível do processo, com cronograma de "
                                 "entregas combinado com a operação.", False, TEXT)]], size=13)
text(s, 0.83, 6.1, 11.6, 0.35, "Energia calculada pelo PCI em base seca do laudo (9.953,8 kcal/kg ≈ 41,7 MJ/kg).",
     size=9, color=GRAY)
footer(s)

# ---------- 10. FICHA TÉCNICA ----------
s = white_slide()
header(s, "Especificação", [("Ficha técnica do ", False), ("produto", True)])
rows = [("Produto", "Resíduo têxtil pós-industrial picotado (mistura de fibras)"),
        ("Aplicação", "Combustível alternativo sólido (substituição parcial)"),
        ("PCI (base seca)", "9.953,8 kcal/kg"),
        ("PCS", "10.040,0 kcal/kg"),
        ("Cloro (base seca)", "< 0,05 %"),
        ("Umidade  /  Sólidos", "3,77 %  /  96,23 %"),
        ("Granulometria", "Picotado, com dimensão ajustável à especificação do cliente"),
        ("Embalagem", "Sacos de até 30 kg"),
        ("Capacidade", "Até 1.000 t/mês, em cargas de 14 t"),
        ("Documentação", "Laudo analítico, MTR e certificado de destinação por lote")]
tbl = s.shapes.add_table(len(rows), 2, Inches(0.83), Inches(1.7), Inches(11.67), Inches(4.8)).table
tbl.columns[0].width = Inches(3.6)
tbl.columns[1].width = Inches(8.07)
for r, (k, v) in enumerate(rows):
    for ci, val in enumerate((k, v)):
        cell = tbl.cell(r, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD if r % 2 == 0 else WHITE
        cell.margin_left = Inches(0.25)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.text = ""
        run = tf.paragraphs[0].add_run()
        run.text = val
        run.font.name = FONT
        run.font.size = Pt(12)
        run.font.bold = ci == 0
        run.font.color.rgb = ROYAL if ci == 0 else TEXT
tbl.first_row = False
tbl.horz_banding = False
footer(s)

# ---------- 11. CONFORMIDADE ----------
s = white_slide()
header(s, "Conformidade", [("Rastreabilidade e ", False), ("conformidade ambiental", True)])
conf = [(ROYAL, "file-text", "MTR", "Manifesto de Transporte de Resíduos emitido em cada coleta e expedição."),
        (ROYAL, "flask-conical", "Laudo analítico", "Caracterização energética (PCI, cloro, umidade) do material fornecido."),
        (ROYAL, "tags", "Controle por lote", "Registro fotográfico, pesagem e identificação de cada lote."),
        (GREEN, "badge-check", "Destinação comprovada", "Certificado de destinação final ambientalmente adequada.")]
for i, (c, ic, t, b) in enumerate(conf):
    card(s, 0.83 + i * 2.97, 1.7, 2.75, 2.9, c, t, b, title_lines=2, ico=ic)
hline(s, 0.83, 4.85, 11.67)
text(s, 0.83, 5.0, 11.6, 1.4, [[("Rastreabilidade completa: ", True, ROYAL), ("da indústria geradora ao forno da mineradora.", False, TEXT)],
                               [("Apoio técnico à mineradora nas etapas de licenciamento para uso de combustível derivado de "
                                 "resíduo (ex.: Resolução SIMA nº 47/2020 no Estado de São Paulo).", False, GRAY, 11)]],
     size=15, line_spacing=1.4)
footer(s)

# ---------- 11b. AMOSTRAS ----------
s = white_slide()
header(s, "Amostras para avaliação", [("Enviaremos ", False), ("15 kg de amostras", True), (" em três formatos", False)],
       sub="5 kg de cada apresentação do material, para análise e teste de queima na sua unidade.")
amostras = [(ROYAL, "scissors", "399170.jpg", "Picotado", "Fragmentos de tecido cortados, prontos para alimentação direta."),
            (BLUE, "layers", "399171.jpg", "Moído", "Fibra triturada fina, homogênea e de fácil dosagem."),
            (GREEN, "mountain", "399172.jpg", "Pedra-brita", "Material compactado em blocos, mais denso para transporte e estocagem.")]
for i, (c, ic, foto, t, b) in enumerate(amostras):
    x = 0.83 + i * 3.95
    img = fit(os.path.join(ARQ, foto), os.path.join(ASSETS, f"amostra_{i + 1}.jpg"), (1000, 560))
    s.shapes.add_picture(img, Inches(x), Inches(2.2), Inches(3.75), Inches(2.1))
    rect(s, x, 4.3, 3.75, 1.85, CARD)
    rect(s, x, 4.3, 3.75, 0.06, c)
    icon(s, ic, x + 0.3, 4.58, 0.4, c)
    text(s, x + 0.85, 4.55, 1.9, 0.45, [[(t, True)]], size=14, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 2.55, 4.5, 0.95, 0.55, [[("5", True, c, 26), (" kg", False, c, 12)]], align=PP_ALIGN.RIGHT)
    text(s, x + 0.3, 5.2, 3.15, 0.8, b, size=10, color=GRAY)
text(s, 0.83, 6.3, 11.6, 0.35, "Amostras identificadas por lote e acompanhadas da ficha técnica do material.",
     size=9, color=GRAY)
footer(s)

# ---------- 12. PRÓXIMOS PASSOS ----------
s = white_slide()
header(s, "Implantação", [("Próximos ", False), ("passos", True)])
nxt = [("send", "Envio de amostras", "5 kg de picotado, 5 kg de moído e 5 kg de pedra-brita."),
       ("flame", "Teste de queima", "Ensaio controlado no equipamento de destino."),
       ("clipboard-check", "Licenciamento", "Adequação da licença para uso do combustível alternativo."),
       ("play", "Fornecimento piloto", "Lotes iniciais com monitoramento de desempenho."),
       ("handshake", "Contrato contínuo", "Fornecimento regular e programado.")]
for i, (ic, t, b) in enumerate(nxt):
    x = 0.83 + i * 2.37
    c = GREEN if i == len(nxt) - 1 else ROYAL
    rect(s, x, 2.0, 2.2, 3.0, CARD)
    rect(s, x, 2.0, 2.2, 0.06, c)
    text(s, x + 0.25, 2.3, 1.5, 0.5, f"{i + 1:02d}", size=24, color=c)
    icon(s, ic, x + 2.2 - 0.65, 2.35, 0.4, c)
    text(s, x + 0.25, 3.0, 1.8, 0.6, [[(t, True)]], size=12)
    text(s, x + 0.25, 3.65, 1.8, 1.2, b, size=10, color=GRAY)
text(s, 0.83, 5.35, 11.6, 0.8, [[("Solicitamos cotação formal em R$/t ", True, ROYAL),
                                ("para o fornecimento de 1.000 t/mês com entrega programada na unidade do cliente. "
                                 "As condições finais serão definidas em reunião.", False, GRAY)]], size=12)
footer(s)

# ---------- 13. ENCERRAMENTO ----------
s = prs.slides.add_slide(BLANK)
dark_bg(s, end_img)
eyebrow(s, 0.83, 1.3, "Nunes & Lucato  ·  Mineração", color=LIGHT)
text(s, 0.83, 1.75, 7.0, 2.6, [[("Transformar o resíduo têxtil em ", False), ("energia para a mineração", True),
                               (" reduz custos, diminui o uso de combustíveis fósseis e fortalece a sua ", False),
                               ("agenda de economia circular", True), (".", False)]], size=22, color=WHITE,
     line_spacing=1.3)
hline(s, 0.83, 4.55, 1.1, LIGHT, 1)
s.shapes.add_picture(os.path.join(ASSETS, "logo_white.png"), Inches(0.83), Inches(4.8), height=Inches(0.55))
text(s, 0.83, 5.65, 6, 0.4, "“Transformando resíduos em valor.”", size=16, color=LIGHT)

# remove o estilo de tema (sombra/contorno padrão) de formas e linhas, para traços finos e limpos
for sl in prs.slides:
    for shp in sl.shapes:
        for st in shp._element.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}style"):
            shp._element.remove(st)

cp = prs.core_properties
cp.author = cp.last_modified_by = "Nunes & Lucato"
cp.title = "Resíduo têxtil como combustível alternativo para a mineração"
cp.subject = "Proposta comercial Nunes & Lucato"
cp.comments = cp.keywords = cp.category = ""

out = os.path.join(OUT_DIR, "Nunes_Lucato_Residuo_Textil_Mineradora.pptx")
prs.save(out)
print("salvo:", out, "slides:", len(prs.slides))
